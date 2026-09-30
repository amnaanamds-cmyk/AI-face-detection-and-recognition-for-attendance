"""Daily / monthly / session reports and export to CSV, Excel and PDF."""
from __future__ import annotations

import calendar
import csv
import io
from dataclasses import dataclass
from datetime import date

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Attendance, AttendanceStatus, ClassSession, Course, Enrollment, Student
from app.services.analytics import rate
from app.terminology import Terms, terms


@dataclass
class Report:
    title: str
    subtitle: str
    headers: list[str]
    rows: list[list]


def _minutes(r) -> int | None:
    if r is None or r.marked_at is None or r.checked_out_at is None:
        return None
    return max(0, int((r.checked_out_at - r.marked_at).total_seconds() // 60))


def _hm(minutes: int | None) -> str:
    return "--" if minutes is None else f"{minutes // 60}:{minutes % 60:02d}"


def daily_report(db: Session, day: date, course_ids: list[int] | None = None, t: Terms | None = None) -> Report:
    t = t or terms(None)
    q = select(ClassSession).where(ClassSession.date == day)
    if course_ids is not None:
        q = q.where(ClassSession.course_id.in_(course_ids))
    rows = []
    for sess in db.scalars(q.order_by(ClassSession.start_time)).all():
        recs = {r.student_id: r for r in sess.attendance}
        students = db.scalars(
            select(Student).join(Enrollment).where(Enrollment.course_id == sess.course_id).order_by(Student.student_code)
        ).all()
        for s in students:
            r = recs.get(s.id)
            row = [
                sess.course.code, sess.start_time.strftime("%H:%M"), s.student_code, s.name,
                r.status.value.title() if r else "Not marked",
                r.marked_at.strftime("%H:%M") if r and r.marked_at else "--",
            ]
            if t.check_out:
                row += [r.checked_out_at.strftime("%H:%M") if r and r.checked_out_at else "--", _hm(_minutes(r))]
            row.append(f"{r.confidence:.2f}" if r and r.confidence is not None else "")
            rows.append(row)
    headers = [t.group, t.session, t.person_id, "Name", "Status", "Check-in" if t.check_out else "Time"]
    if t.check_out:
        headers += ["Check-out", "Hours"]
    return Report("Daily Attendance Report", f"Date: {day:%d %B %Y}", headers + ["Confidence"], rows)


def session_report(db: Session, session: ClassSession, t: Terms | None = None) -> Report:
    t = t or terms(None)
    r = daily_report(db, session.date, [session.course_id], t)
    start = session.start_time.strftime("%H:%M")
    r.rows = [row for row in r.rows if row[1] == start]
    r.title = f"{t.session} Report - {session.course.code} {session.course.name}"
    r.subtitle = f"{session.date:%d %B %Y}, {start}-{session.end_time:%H:%M}"
    return r


def monthly_report(db: Session, year: int, month: int, course_ids: list[int] | None = None,
                   t: Terms | None = None) -> Report:
    t = t or terms(None)
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])
    q = select(Attendance).where(Attendance.date >= start, Attendance.date <= end)
    if course_ids is not None:
        q = q.where(Attendance.course_id.in_(course_ids))
    per_student: dict[int, list[AttendanceStatus]] = {}
    minutes: dict[int, int] = {}
    for r in db.scalars(q).all():
        per_student.setdefault(r.student_id, []).append(r.status)
        minutes[r.student_id] = minutes.get(r.student_id, 0) + (_minutes(r) or 0)
    students = {s.id: s for s in db.scalars(select(Student).where(Student.id.in_(list(per_student)))).all()}
    rows = []
    for sid in sorted(per_student, key=lambda i: students[i].student_code):
        st = per_student[sid]
        pct = rate(st)
        row = [
            students[sid].student_code, students[sid].name,
            st.count(AttendanceStatus.present), st.count(AttendanceStatus.late), st.count(AttendanceStatus.absent),
            st.count(AttendanceStatus.excused) + st.count(AttendanceStatus.leave),
            f"{pct:.1f}%" if pct is not None else "--",
        ]
        if t.check_out:
            row.append(_hm(minutes.get(sid)))
        rows.append(row)
    scope = f"All {t.groups.lower()}"
    if course_ids is not None:
        scope = ", ".join(c.code for c in db.scalars(select(Course).where(Course.id.in_(course_ids))).all()) or "-"
    headers = [t.person_id, "Name", "Present", "Late", "Absent", "Excused/Leave", "Percentage"]
    if t.check_out:
        headers.append("Hours on site")
    return Report("Monthly Attendance Report", f"{start:%B %Y} - {scope}", headers, rows)


# ------------------------------------------------------------------------ exporters
def to_csv(report: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([report.title])
    w.writerow([report.subtitle])
    w.writerow([])
    w.writerow(report.headers)
    w.writerows(report.rows)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([report.subtitle])
    ws.append([])
    ws.append(report.headers)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for row in report.rows:
        ws.append(row)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col[3:]) if len(col) > 3 else 10
        ws.column_dimensions[col[0].column_letter].width = min(40, max(10, width + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=report.title)
    styles = getSampleStyleSheet()
    story = [Paragraph(report.title, styles["Title"]), Paragraph(report.subtitle, styles["Normal"]), Spacer(1, 12)]
    data = [report.headers] + [[str(c) for c in row] for row in report.rows]
    if not report.rows:
        data.append(["No records"] + [""] * (len(report.headers) - 1))
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    story.append(table)
    doc.build(story)
    return buf.getvalue()


EXPORTERS = {
    "csv": (to_csv, "text/csv"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": (to_pdf, "application/pdf"),
}
