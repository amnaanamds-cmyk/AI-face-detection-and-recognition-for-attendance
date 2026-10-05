"""School overview for the principal: one page that combines every teacher's attendance.

Teachers take attendance on their own phones (each with their own login); every record goes to the
same server, so the principal sees the whole school at once:

* today: which teachers have taken attendance, with present / expected per class;
* teachers: classes held, the attendance % of their students, how much was marked by hand;
* subjects: attendance % per subject and how many students are below the required %;
* students: overall % plus a % per subject (the register a principal signs off).

Percentages follow app/services/analytics.rate (Excused and Leave are neutral).
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (Attendance, AttendanceStatus, ClassSession, Course, Enrollment, Role, SessionState, Student,
                        User)
from app.services.analytics import ATTENDED, rate

PERIODS = {"today": "Today", "week": "This week", "month": "This month", "all": "Whole term"}


def period_range(period: str, today: date | None = None, start: str = "", end: str = "") -> tuple[date | None, date]:
    """(first day, last day) of a period; first day None = since the beginning."""
    today = today or date.today()
    if start or end:
        try:
            s = date.fromisoformat(start) if start else None
            e = date.fromisoformat(end) if end else today
        except ValueError:
            return today, today
        return (min(s, e), max(s, e)) if s else (None, e)
    if period == "week":
        return today - timedelta(days=today.weekday()), today
    if period == "month":
        return today.replace(day=1), today
    if period == "all":
        return None, today
    return today, today


def _class_label(semester, section) -> str:
    return f"{semester}-{section}" if section else str(semester)


def overview(db: Session, org_id: int, start: date | None, end: date, threshold: float = 75.0,
             today: date | None = None) -> dict:
    today = today or date.today()
    courses = db.scalars(select(Course).where(Course.org_id == org_id).order_by(Course.code)).all()
    course_ids = [c.id for c in courses]

    sq = select(ClassSession).where(ClassSession.course_id.in_(course_ids), ClassSession.date <= end)
    rq = select(Attendance).where(Attendance.course_id.in_(course_ids), Attendance.date <= end)
    if start:
        sq = sq.where(ClassSession.date >= start)
        rq = rq.where(Attendance.date >= start)
    sessions = [s for s in db.scalars(sq).all() if s.state != SessionState.scheduled]
    records = db.scalars(rq).all()

    enrolled: dict[int, list[int]] = defaultdict(list)          # course -> students
    for course_id, student_id in db.execute(select(Enrollment.course_id, Enrollment.student_id)
                                            .where(Enrollment.course_id.in_(course_ids))).all():
        enrolled[course_id].append(student_id)

    by_course = defaultdict(list)
    by_student = defaultdict(list)
    by_student_course = defaultdict(list)
    for r in records:
        by_course[r.course_id].append(r)
        by_student[r.student_id].append(r.status)
        by_student_course[(r.student_id, r.course_id)].append(r.status)
    for (sid, cid) in by_student_course:                         # records without an enrollment still count
        if sid not in enrolled[cid]:
            enrolled[cid].append(sid)
    held = defaultdict(list)
    for s in sessions:
        held[s.course_id].append(s)

    # --- subjects -----------------------------------------------------------------------------------
    subjects = []
    for c in courses:
        if c.code == "GENERAL" and not by_course.get(c.id):
            continue
        st_rates = [rate(by_student_course.get((sid, c.id), [])) for sid in enrolled.get(c.id, [])]
        subjects.append({
            "course": c, "class": _class_label(c.semester, c.section),
            "teacher": c.teacher.full_name if c.teacher else "-",
            "held": len(held.get(c.id, [])), "students": len(enrolled.get(c.id, [])),
            "rate": rate(r.status for r in by_course.get(c.id, [])),
            "below": sum(1 for x in st_rates if x is not None and x < threshold),
        })

    # --- teachers -----------------------------------------------------------------------------------
    teachers = []
    today_sessions = defaultdict(list)
    for s in db.scalars(select(ClassSession).where(ClassSession.course_id.in_(course_ids),
                                                   ClassSession.date == today)).all():
        today_sessions[s.course_id].append(s)
    today_recs = defaultdict(list)
    for r in db.scalars(select(Attendance).where(Attendance.course_id.in_(course_ids), Attendance.date == today)).all():
        today_recs[r.session_id].append(r.status)

    for u in db.scalars(select(User).where(User.org_id == org_id, User.role == Role.teacher, User.is_active.is_(True))
                        .order_by(User.full_name)).all():
        mine = [c for c in courses if c.teacher_id == u.id]
        recs = [r for c in mine for r in by_course.get(c.id, [])]
        marked = [r for r in recs if r.status in ATTENDED]
        my_sessions = [s for c in mine for s in held.get(c.id, [])]
        last = max((s.start_time for c in mine for s in db.scalars(
            select(ClassSession).where(ClassSession.course_id == c.id, ClassSession.state != SessionState.scheduled)
            .order_by(ClassSession.start_time.desc()).limit(1))), default=None)
        now_classes = []
        for c in mine:
            for s in today_sessions.get(c.id, []):
                st = today_recs.get(s.id, [])
                now_classes.append({"session": s, "course": c, "expected": len(enrolled.get(c.id, [])),
                                    "present": sum(1 for x in st if x in ATTENDED)})
        teachers.append({
            "user": u, "subjects": mine, "held": len(my_sessions), "rate": rate(r.status for r in recs),
            "by_hand": round(100.0 * sum(1 for r in marked if r.method not in ("face", "kiosk")) / len(marked), 1)
            if marked else None,
            "last": last, "today": now_classes,
            "taken_today": any(x["session"].state != SessionState.scheduled for x in now_classes),
        })

    # --- students -----------------------------------------------------------------------------------
    student_courses = defaultdict(list)
    for cid, sids in enrolled.items():
        for sid in sids:
            student_courses[sid].append(cid)
    students = []
    for s in db.scalars(select(Student).where(Student.org_id == org_id, Student.is_active.is_(True))
                        .order_by(Student.semester, Student.section, Student.student_code)).all():
        st = by_student.get(s.id, [])
        students.append({
            "student": s, "class": _class_label(s.semester, s.section), "rate": rate(st),
            "attended": sum(1 for x in st if x in ATTENDED),
            "counted": sum(1 for x in st if x not in (AttendanceStatus.excused, AttendanceStatus.leave)),
            "absent": st.count(AttendanceStatus.absent),
            "per_subject": {cid: rate(by_student_course.get((s.id, cid), [])) for cid in student_courses.get(s.id, [])},
        })

    rated = [x for x in students if x["rate"] is not None]
    teachers_with_subjects = [t for t in teachers if t["subjects"]]
    return {
        "start": start, "end": end, "threshold": threshold,
        "rate": rate(r.status for r in records), "records": len(records), "held": len(sessions),
        "below": sum(1 for x in rated if x["rate"] < threshold), "rated": len(rated),
        "teachers_total": len(teachers_with_subjects),
        "teachers_taken": sum(1 for t in teachers_with_subjects if t["taken_today"]),
        "subjects": subjects, "teachers": teachers, "students": students,
        "subject_cols": [x["course"] for x in subjects],
        "classes": sorted({x["class"] for x in students}),
    }


def to_xlsx(data: dict, org_name: str = "") -> bytes:
    """The overview as an Excel workbook: Students (with a column per subject), Subjects, Teachers."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    period = f"{data['start'] or 'start'} to {data['end']}"
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    low = PatternFill("solid", fgColor="F8D7DA")

    def sheet(ws, title, header, rows, rate_cols=()):
        ws.title = title
        ws.append([f"{org_name} - {title} attendance, {period}".strip(" -")])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(header)
        for cell in ws[2]:
            cell.font, cell.fill = head, fill
        for row in rows:
            ws.append(row)
            for i in rate_cols:
                v = row[i]
                if isinstance(v, (int, float)) and v < data["threshold"]:
                    ws.cell(ws.max_row, i + 1).fill = low
        for col in ws.columns:
            ws.column_dimensions[col[1].column_letter].width = max(10, min(40, max(len(str(c.value or "")) for c in col[1:]) + 2))

    cols = data["subject_cols"]
    sheet(wb.active, "Students", ["ID", "Name", "Class", "Attended", "Counted", "Absent", "Overall %"] + [c.code for c in cols],
          [[x["student"].student_code, x["student"].name, x["class"], x["attended"], x["counted"], x["absent"], x["rate"]]
           + [x["per_subject"].get(c.id) for c in cols] for x in data["students"]],
          rate_cols=range(6, 7 + len(cols)))
    sheet(wb.create_sheet(), "Subjects", ["Code", "Subject", "Class", "Teacher", "Classes held", "Students", "Attendance %",
                                          f"Students below {data['threshold']:.0f}%"],
          [[x["course"].code, x["course"].name, x["class"], x["teacher"], x["held"], x["students"], x["rate"], x["below"]]
           for x in data["subjects"]], rate_cols=(6,))
    sheet(wb.create_sheet(), "Teachers", ["Teacher", "Subjects", "Classes held", "Students' attendance %",
                                          "Marked by hand %", "Last attendance taken"],
          [[x["user"].full_name, ", ".join(c.code for c in x["subjects"]), x["held"], x["rate"], x["by_hand"],
            x["last"].strftime("%Y-%m-%d %H:%M") if x["last"] else "never"] for x in data["teachers"]], rate_cols=(3,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
