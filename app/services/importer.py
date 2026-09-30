"""Bulk import of real class lists (CSV / Excel) and face photos (ZIP / folder)."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Course, Enrollment, Student, now
from app.services import faces
from app.vision.base import decode_image

# accepted column names (lower-case, spaces/underscores/dots removed) -> Student attribute
COLUMN_ALIASES = {
    "studentid": "student_code", "studentcode": "student_code", "id": "student_code", "regno": "student_code",
    "registrationno": "student_code", "registrationnumber": "student_code", "enrollmentno": "student_code",
    "name": "name", "studentname": "name", "fullname": "name",
    "rollno": "roll_number", "rollnumber": "roll_number", "roll": "roll_number",
    "department": "department", "dept": "department", "program": "department", "programme": "department",
    "semester": "semester", "sem": "semester",
    "section": "section", "sec": "section",
    "email": "email", "emailaddress": "email",
    "phone": "phone", "contact": "phone", "mobile": "phone", "contactno": "phone",
    "consent": "consent_given", "consentgiven": "consent_given", "biometricconsent": "consent_given",
}
IMG_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}


@dataclass
class ImportReport:
    created: int = 0
    updated: int = 0
    enrolled_courses: int = 0
    errors: list[str] = field(default_factory=list)


def _norm(col: str) -> str:
    return re.sub(r"[\s_.\-#/]", "", str(col or "").strip().lower())


def read_table(filename: str, data: bytes) -> list[dict[str, str]]:
    """Read a CSV or Excel file into a list of {original column: value} dicts."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            return []
        header = rows[0]
        return [dict(zip(header, r)) for r in rows[1:]]
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
    return [{k: (v or "").strip() for k, v in row.items() if k} for row in csv.DictReader(io.StringIO(text), dialect=dialect)]


def _truthy(v: str) -> bool:
    return str(v).strip().lower() in {"1", "yes", "y", "true", "x", "✓"}


def import_students(db: Session, org_id: int, rows: list[dict[str, str]], *, default_consent: bool = False,
                    auto_enroll: bool = True) -> ImportReport:
    rep = ImportReport()
    if rows:
        unknown = [c for c in rows[0] if _norm(c) not in COLUMN_ALIASES]
        mapped = {COLUMN_ALIASES[_norm(c)] for c in rows[0] if _norm(c) in COLUMN_ALIASES}
        if not {"student_code", "name"} <= mapped:
            rep.errors.append("file needs at least a 'Student ID' and a 'Name' column "
                              f"(found: {', '.join(rows[0])})")
            return rep
        if unknown:
            rep.errors.append(f"ignored columns: {', '.join(unknown)}")
    from app.models import Organization
    from app.services.billing import plan_limit_error

    org = db.get(Organization, org_id)
    courses = db.scalars(select(Course).where(Course.org_id == org_id)).all()
    for line, raw in enumerate(rows, start=2):
        data = {COLUMN_ALIASES[_norm(k)]: v for k, v in raw.items() if _norm(k) in COLUMN_ALIASES}
        code, name = data.get("student_code", "").strip(), data.get("name", "").strip()
        if not code or not name:
            rep.errors.append(f"row {line}: missing Student ID or name - skipped")
            continue
        try:
            semester = int(float(data["semester"])) if data.get("semester") else None
        except ValueError:
            rep.errors.append(f"row {line}: invalid semester '{data['semester']}' - skipped")
            continue
        st = db.scalar(select(Student).where(Student.org_id == org_id, Student.student_code == code))
        if st is None:
            limit = plan_limit_error(db, org)
            if limit:
                rep.errors.append(f"row {line}: {limit} - stopped here")
                break
            st = Student(student_code=code, name=name, org_id=org_id)
            db.add(st)
            rep.created += 1
        else:
            rep.updated += 1
        st.name = name
        for attr in ("roll_number", "department", "section", "email", "phone"):
            if data.get(attr):
                setattr(st, attr, data[attr])
        if semester is not None:
            st.semester = semester
        was = st.consent_given
        if "consent_given" in data:
            st.consent_given = _truthy(data["consent_given"])
        elif default_consent and not st.consent_given:
            st.consent_given = True
        if st.consent_given and not was:
            st.consent_at = now()
        st.department = st.department or "Computer Science"
        st.section = st.section or "A"
        st.semester = st.semester or 1
        db.flush()
        if auto_enroll:
            have = {e.course_id for e in db.scalars(select(Enrollment).where(Enrollment.student_id == st.id))}
            for c in courses:
                if (c.department, c.semester, c.section) == (st.department, st.semester, st.section) and c.id not in have:
                    db.add(Enrollment(student_id=st.id, course_id=c.id))
                    rep.enrolled_courses += 1
    db.commit()
    faces.gallery_cache.invalidate(org_id)
    return rep


def _student_code_from_path(path: str) -> str | None:
    """'BSCS-2023-001/front.jpg' or 'BSCS-2023-001_2.jpg' or 'BSCS-2023-001.jpg' -> 'BSCS-2023-001'."""
    p = Path(path)
    if p.suffix.lower() not in IMG_EXT or p.name.startswith(".") or "__MACOSX" in p.parts:
        return None
    if len(p.parts) >= 2:
        return p.parts[-2]
    stem = p.stem
    m = re.match(r"^(.*?)[_ ](\d+|front|left|right|up|down)$", stem, re.IGNORECASE)
    return m.group(1) if m else stem


@dataclass
class FaceImportReport:
    students: int = 0
    templates: int = 0
    messages: list[str] = field(default_factory=list)


def import_faces(db: Session, org_id: int, files: list[tuple[str, bytes]]) -> FaceImportReport:
    """Enroll faces from (path, bytes) pairs grouped by the student ID in the path."""
    groups: dict[str, list[bytes]] = defaultdict(list)
    for path, data in files:
        code = _student_code_from_path(path)
        if code:
            groups[code].append(data)
    rep = FaceImportReport()
    for code, blobs in sorted(groups.items()):
        st = db.scalar(select(Student).where(Student.org_id == org_id, Student.student_code == code))
        if st is None:
            rep.messages.append(f"{code}: no student with this ID - {len(blobs)} photo(s) skipped")
            continue
        if not st.consent_given:
            rep.messages.append(f"{code}: biometric consent not recorded - skipped")
            continue
        images: list[np.ndarray] = []
        for b in blobs:
            try:
                images.append(decode_image(b))
            except ValueError:
                rep.messages.append(f"{code}: unreadable image skipped")
        r = faces.enroll_images(db, st, images)
        if r.accepted:
            rep.students += 1
            rep.templates += r.accepted
        rep.messages.append(f"{code}: {r.accepted} template(s) added"
                            + (f"; {'; '.join(r.rejected)}" if r.rejected else ""))
    return rep


def read_zip(data: bytes) -> list[tuple[str, bytes]]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return [(i.filename, zf.read(i)) for i in zf.infolist()
                if not i.is_dir() and Path(i.filename).suffix.lower() in IMG_EXT and i.file_size < 20_000_000]


def read_folder(folder: Path) -> list[tuple[str, bytes]]:
    return [(str(p.relative_to(folder)), p.read_bytes()) for p in sorted(folder.rglob("*"))
            if p.is_file() and p.suffix.lower() in IMG_EXT]
