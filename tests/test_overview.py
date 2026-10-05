"""Principal's school overview: every teacher's attendance combined, with % per student, subject and teacher."""
import io
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from app import database
from app.models import Attendance, AttendanceStatus as S, ClassSession, Course, SessionState, Student
from app.services import attendance as att
from app.services import overview as ov
from tests.conftest import login


def school(client):
    """Two teachers on their own phones, two subjects, three students."""
    login(client)
    for u, n in (("t.math", "Ms. Math"), ("t.eng", "Mr. English")):
        client.post("/users/new", data={"username": u, "full_name": n, "password": "teacher123", "role": "teacher"})
    client.post("/courses/new", data={"code": "MATH", "name": "Mathematics", "semester": 9, "section": "A", "teacher_id": 2})
    client.post("/courses/new", data={"code": "ENG", "name": "English", "semester": 9, "section": "A", "teacher_id": 3})
    for i, name in enumerate(["Amina", "Bilal", "Sana"], start=1):
        client.post("/students/new", data={"student_code": f"S-{i}", "name": name, "semester": 9, "section": "A",
                                           "consent": "on", "auto_enroll": "on"})


def record(db, course_code, day, statuses, method="face", state=SessionState.closed):
    c = db.query(Course).filter_by(code=course_code).one()
    s = ClassSession(course_id=c.id, date=day, start_time=datetime.combine(day, datetime.min.time()) + timedelta(hours=9),
                     state=state, created_by=c.teacher_id)
    db.add(s)
    db.flush()
    for code, st in statuses.items():
        stu = db.query(Student).filter_by(student_code=code).one()
        db.add(Attendance(student_id=stu.id, session_id=s.id, course_id=c.id, date=day, status=st, method=method,
                          marked_at=s.start_time if st != S.absent else None))
    db.commit()
    return s


def test_percentages_per_student_subject_teacher(client):
    school(client)
    today = date.today()
    db = database.SessionLocal()
    record(db, "MATH", today - timedelta(days=1), {"S-1": S.present, "S-2": S.absent, "S-3": S.present})
    record(db, "MATH", today, {"S-1": S.present, "S-2": S.absent, "S-3": S.late})
    record(db, "ENG", today - timedelta(days=1), {"S-1": S.present, "S-2": S.present, "S-3": S.leave}, method="manual")
    org_id = db.query(Course).first().org_id

    d = ov.overview(db, org_id, None, today, threshold=75)
    assert d["rate"] == 75.0 and d["held"] == 3 and d["records"] == 9
    subj = {x["course"].code: x for x in d["subjects"]}
    assert subj["MATH"]["rate"] == round(100 * 4 / 6, 1) and subj["MATH"]["below"] == 1 and subj["MATH"]["teacher"] == "Ms. Math"
    assert subj["ENG"]["rate"] == 100.0 and subj["ENG"]["held"] == 1
    stu = {x["student"].student_code: x for x in d["students"]}
    assert stu["S-2"]["rate"] == round(100 / 3, 1) and stu["S-2"]["absent"] == 2
    assert stu["S-3"]["per_subject"][subj["ENG"]["course"].id] is None          # only leave: neutral
    assert stu["S-1"]["per_subject"][subj["MATH"]["course"].id] == 100.0
    assert d["below"] == 1
    tea = {x["user"].username: x for x in d["teachers"]}
    assert tea["t.math"]["held"] == 2 and tea["t.math"]["taken_today"] and tea["t.math"]["by_hand"] == 0.0
    assert not tea["t.eng"]["taken_today"] and tea["t.eng"]["by_hand"] == 100.0
    assert (d["teachers_taken"], d["teachers_total"]) == (1, 2)

    today_only = ov.overview(db, org_id, today, today)
    assert today_only["records"] == 3 and today_only["held"] == 1
    db.close()


def test_period_range():
    d = date(2026, 10, 7)  # a Wednesday
    assert ov.period_range("today", d) == (d, d)
    assert ov.period_range("week", d) == (date(2026, 10, 5), d)
    assert ov.period_range("month", d) == (date(2026, 10, 1), d)
    assert ov.period_range("all", d) == (None, d)
    assert ov.period_range("", d, "2026-10-09", "2026-10-01") == (date(2026, 10, 1), date(2026, 10, 9))
    assert ov.period_range("", d, "bad") == (d, d)


def test_page_excel_and_access(client):
    school(client)
    db = database.SessionLocal()
    record(db, "MATH", date.today(), {"S-1": S.present, "S-2": S.absent})
    db.close()
    r = client.get("/overview")
    assert r.status_code == 200 and "Ms. Math" in r.text and "Mathematics" in r.text and "Amina" in r.text
    assert client.get("/overview?period=all").status_code == 200
    assert client.get("/overview?start=2026-01-01&end=2026-12-31").status_code == 200
    client.post("/settings", data={"org_name": "Govt. High School Model Town"})
    assert "Govt. High School Model Town" in client.get("/overview").text
    x = client.get("/overview.xlsx?period=month")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Students", "Subjects", "Teachers"]
    assert [c.value for c in wb["Students"][2]][:7] == ["ID", "Name", "Class", "Attended", "Counted", "Absent", "Overall %"]
    client.get("/logout")
    login(client, "t.math", "teacher123")           # a teacher sees only their own subjects, not the school
    assert client.get("/overview").status_code == 403
    assert client.get("/overview.xlsx").status_code == 403


def test_forgotten_session_is_closed_and_absentees_saved(client):
    school(client)
    db = database.SessionLocal()
    day = date.today() - timedelta(days=1)
    s = record(db, "MATH", day, {"S-1": S.present}, state=SessionState.active)
    org_id = s.course.org_id
    att.sync_scheduled_sessions(db, org_id, when=s.end_time + timedelta(minutes=30))
    db.refresh(s)
    assert s.state == SessionState.active                          # teacher may still be using it
    att.sync_scheduled_sessions(db, org_id, when=s.end_time + att.FORGOTTEN_GRACE)
    db.refresh(s)
    assert s.state == SessionState.closed
    saved = {a.student.student_code: a.status for a in db.query(Attendance).filter_by(session_id=s.id)}
    assert saved == {"S-1": S.present, "S-2": S.absent, "S-3": S.absent}
    db.close()
